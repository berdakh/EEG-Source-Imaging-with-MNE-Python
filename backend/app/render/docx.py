"""Write light Markdown (headings, bullets, paragraphs) into a .docx, optionally on a template."""

from __future__ import annotations

from pathlib import Path

import docx


def write_docx(text: str, path: Path, template: Path | None = None) -> Path:
    doc = docx.Document(str(template)) if template and template.exists() else docx.Document()
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.lstrip().startswith(("- ", "* ")):
            doc.add_paragraph(line.lstrip()[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
